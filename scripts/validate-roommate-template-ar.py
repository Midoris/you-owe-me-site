#!/usr/bin/env python3
"""Read-only preservation checks for the reviewed localized roommate files."""
import csv
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
COPY = json.loads((ROOT / "content/roommate-workbooks/ar.json").read_text())
NAMES = COPY["sheet_names"]
STRINGS = COPY["strings"]


def translated(value):
    if not isinstance(value, str):
        return value
    if value.startswith("="):
        for original, localized in sorted(STRINGS.items(), key=lambda item: -len(item[0])):
            value = value.replace('"' + original + '"', '"' + localized + '"')
        for original, localized in NAMES.items():
            value = value.replace("'" + original + "'!", "'" + localized + "'!")
        return value
    return STRINGS.get(value, value)


def validate_common(path):
    book = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    assert book.sheetnames == list(NAMES.values())
    assert all(sheet.sheet_view.rightToLeft is True for sheet in book)
    text_formulas = [cell for sheet in book for row in sheet for cell in row if cell.data_type == "f" and cell.number_format == "General"]
    assert len(text_formulas) == (455 if "Google Sheets" in values["ابدأ هنا"]["A1"].value else 453)
    assert all(cell.alignment.readingOrder == 2 for cell in text_formulas)
    assert all(cell.alignment.readingOrder == 0 for sheet in book for row in sheet for cell in row if isinstance(cell.value, (int, float)) or cell.number_format != "General")
    assert [values["القوائم"].cell(row,4).value for row in range(5,17)] == ["USD","EUR","GBP","CAD","AUD","NZD","JPY","CHF","SGD","THB","SAR","أخرى"]
    assert values["القوائم"]["A12"].value == "متفرقات"
    assert all(values["القوائم"].cell(16,col).value in (None, "") for col in range(1,4))
    assert all(values["الإعداد"].cell(row, 2).value in (None, "") for row in range(9, 15))
    assert [values["الإعداد"].cell(row, 3).value for row in range(9, 15)] == [0] * 6
    assert values["القوائم"]["D15"].value == "SAR"
    assert values["القوائم"]["D16"].value == "أخرى"
    assert [values["مثال"].cell(row, 1).value for row in range(16, 19)] == ["سارة", "عمر", "ليلى"]
    assert [values["مثال"].cell(row, 6).value for row in range(16, 19)] == [940, -400, -540]
    assert values["مثال"]["B20"].value == 1980
    assert values["مثال"]["B21"].value == 660
    assert [values["مثال"].cell(row, 3).value for row in (25, 26)] == [400, 540]
    assert values["مثال"]["F24"].value == 0
    for sheet_name in ("المصاريف", "الدفعات"):
        assert book[sheet_name].max_row == 204
        assert book[sheet_name]["A5"].number_format == "dd/mm/yyyy"
    for row in (6, 7, 8, 12):
        assert book["مثال"].cell(row, 1).number_format == "dd/mm/yyyy"
    for row, date in [(6, "2026-05-01"), (7, "2026-05-05"), (8, "2026-05-12"), (12, "2026-05-20")]:
        assert values["مثال"].cell(row, 1).value.date().isoformat() == date
    assert len(book.defined_names) == 5
    assert book["المصاريف"].auto_filter.ref == "A4:Z204"
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "e", (sheet.title, cell.coordinate, cell.value)
    return book, values


def validate_excel():
    path = ROOT / "downloads/ar/roommate-expense-tracker-template.xlsx"
    book, values = validate_common(path)
    original = load_workbook(ROOT / "downloads/roommate-expense-tracker-template.xlsx", data_only=False)
    compared_cells = compared_formulas = 0
    for old in original:
        current = book[NAMES[old.title]]
        assert current.max_row == old.max_row + (old.title == "Settings") and current.max_column == old.max_column
        assert current.sheet_view.rightToLeft is True
        assert str(current.merged_cells) == str(old.merged_cells)
        assert current.freeze_panes == old.freeze_panes
        assert set(current.tables) == set(old.tables)
        assert len(current.data_validations.dataValidation) == len(old.data_validations.dataValidation)
        for before, after in zip(old.data_validations.dataValidation, current.data_validations.dataValidation):
            assert before.sqref == after.sqref
            assert before.type == after.type and before.formula1 == after.formula1
        for row in old:
            for cell in row:
                actual = current[cell.coordinate]
                expected = "SAR" if old.title == "Settings" and cell.coordinate == "D15" else COPY["cell_overrides"].get(f"{old.title}!{cell.coordinate}", translated(cell.value))
                assert actual.value == expected, (old.title, cell.coordinate, cell.value, expected, actual.value)
                compared_cells += 1
                compared_formulas += cell.data_type == "f"
                expected_format = "dd/mm/yyyy" if cell.number_format == "yyyy-mm-dd" else cell.number_format
                assert actual.number_format == expected_format
        for key, dimensions in old.column_dimensions.items():
            assert current.column_dimensions[key].hidden == dimensions.hidden
            assert current.column_dimensions[key].width == dimensions.width
        assert len(old.conditional_formatting) == len(current.conditional_formatting)
        for old_range, new_range in zip(old.conditional_formatting, current.conditional_formatting):
            assert old_range.sqref == new_range.sqref
            old_rules = old.conditional_formatting[old_range]
            new_rules = current.conditional_formatting[new_range]
            assert len(old_rules) == len(new_rules)
            for before, after in zip(old_rules, new_rules):
                assert after.type == before.type
                assert after.text == STRINGS.get(before.text, before.text)
                assert after.formula == [translated("=" + formula)[1:] for formula in before.formula]
    for name, definition in original.defined_names.items():
        expected = definition.attr_text
        for original_name, localized in NAMES.items():
            expected = expected.replace("'" + original_name + "'!", "'" + localized + "'!")
        expected = expected.replace("$D$5:$D$15", "$D$5:$D$16")
        assert book.defined_names[name].attr_text == expected
    with (ROOT / "downloads/ar/roommate-expense-tracker-template.csv").open(newline="", encoding="utf-8") as file:
        rows = list(csv.reader(file))
    assert rows == COPY["csv"] and all(len(row) == 11 for row in rows)
    return {"compared_cells": compared_cells, "compared_formulas": compared_formulas, "native_features_preserved": True, "csv_rows": len(rows)}


def validate_sheets(path):
    book, values = validate_common(path)
    assert not book["الدفعات"].tables
    assert STRINGS["Roommate Expense Tracker — Google Sheets"] == values["ابدأ هنا"]["A1"].value
    assert 'HYPERLINK("https://you-owe-me.com/ar/solutions/app-to-track-money-owed/","تابع المبالغ المتبقية على iPhone")' in book["ابدأ هنا"]["A23"].value
    assert values["ابدأ هنا"]["A24"].value == STRINGS["Start a separate record in You Owe Me; this spreadsheet is not imported."]
    setup = book["الإعداد"]["A17"].value
    for fragment in ('TRIM', 'SUBSTITUTE(B9,"  "," ")', 'NOT(ISNUMBER(C9))', 'ROUND(C16,2)<>0', STRINGS["Ready: setup inputs are valid."]):
        assert fragment in setup
    for row in range(5, 205):
        share = book["المصاريف"].cell(row, 20).value
        check = book["المصاريف"].cell(row, 26).value
        assert f'QUOTIENT(ROUND($D{row}*100,0),COUNTIF($G{row}:$L{row},"نعم"))' in share
        assert f'MOD(ROUND($D{row}*100,0),COUNTIF($G{row}:$L{row},"نعم"))' in share
        assert "SUBSTITUTE" in check and "ISNUMBER" in check
        assert STRINGS["Custom shares must add up to the expense amount."] in check
        assert "SUBSTITUTE" in book["الدفعات"].cell(row, 6).value
    assert "القوائم'!$G$5" in book["الملخص"]["B17"].value
    assert "القوائم'!$G$5" in book["الملخص"]["G5"].value
    assert "ROUND(G5,2)=0" in book["الملخص"]["H5"].value
    assert "<0" in book["الملخص"]["K22"].value
    assert "COUNTIFS('المصاريف'" in book["القوائم"]["G5"].value
    return {"direct_range_edition": True, "readiness_gate": True, "integer_cent_allocation": True, "native_engine_pending": True}


if __name__ == "__main__":
    result = {"excel": validate_excel()}
    if len(sys.argv) == 2:
        result["sheets_intermediary"] = validate_sheets(Path(sys.argv[1]))
    print(json.dumps(result, ensure_ascii=False))
