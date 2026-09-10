#!/usr/bin/env python3
"""Validate the focused Google Sheets import workbook contract without recalculating it."""

from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
import sys
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHEETS = ["Start Here", "Setup", "Expenses", "Repayments", "Summary", "Example", "Settings"]


def formula(cell) -> str:
    value = cell.value
    assert isinstance(value, str) and value.startswith("="), cell.coordinate
    return value


def allocate_equal_shares(amount: str, included: list[bool]) -> list[Decimal]:
    """Mirror the integer-cent contract used by the direct Sheets formula."""
    total_cents = int((Decimal(amount) * 100).to_integral_value(rounding=ROUND_HALF_UP))
    included_count = sum(included)
    quotient, remainder = divmod(total_cents, included_count)
    ordinal = 0
    shares: list[Decimal] = []
    for is_included in included:
        if not is_included:
            shares.append(Decimal("0"))
            continue
        ordinal += 1
        cents = quotient + (1 if ordinal > included_count - remainder else 0)
        shares.append(Decimal(cents) / 100)
    return shares


def validate_equal_share_invariants() -> None:
    cases = [
        ("10.00", [True, True, True], ["3.33", "3.33", "3.34"]),
        ("10.01", [True, True, True], ["3.33", "3.34", "3.34"]),
        ("0.03", [True, True, True, True, True, True], ["0", "0", "0", "0.01", "0.01", "0.01"]),
        ("0.03", [True, False, True, False, True, False], ["0.01", "0", "0.01", "0", "0.01", "0"]),
    ]
    for amount, included, expected in cases:
        shares = allocate_equal_shares(amount, included)
        assert shares == [Decimal(value) for value in expected], (amount, included, shares)
        included_shares = [share for share, is_included in zip(shares, included) if is_included]
        assert sum(shares) == Decimal(amount)
        assert min(included_shares) >= 0
        assert max(included_shares) - min(included_shares) <= Decimal("0.01")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: validate-roommate-google-sheets-import.py <google-sheets-import.xlsx>")

    path = Path(sys.argv[1]).resolve()
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)

    assert formulas.sheetnames == EXPECTED_SHEETS
    assert values["Start Here"]["A1"].value == "Roommate Expense Tracker — Google Sheets"
    for required_text in [
        "Make your own copy in Google Drive before entering personal information.",
        "Your copy is stored in your Google account. Choose who can view or edit it. You Owe Me does not receive the expenses you enter.",
        "Track roommate balances on iPhone",
        "Start a separate record in You Owe Me; this spreadsheet is not imported.",
    ]:
        assert any(required_text in str(cell.value or "") for row in values["Start Here"].iter_rows() for cell in row), required_text

    setup_values = values["Setup"]
    setup_formulas = formulas["Setup"]
    assert all(setup_values.cell(row, 2).value in (None, "") for row in range(9, 15))
    assert [setup_values.cell(row, 3).value for row in range(9, 15)] == [0] * 6
    setup_check = formula(setup_formulas["A17"])
    for fragment in [
        "TRIM",
        'SUBSTITUTE(B9,"  "," ")',
        "NOT(ISNUMBER(C9))",
        "Set opening positions for unused roommate slots to 0.00.",
        "ROUND(C16,2)<>0",
        "Ready: setup inputs are valid.",
    ]:
        assert fragment in setup_check

    start_link = formula(formulas["Start Here"]["A23"])
    assert 'HYPERLINK("https://you-owe-me.com/solutions/roommate-expense-tracker/","Track roommate balances on iPhone")' in start_link

    expenses = formulas["Expenses"]
    # The import edition deliberately avoids Excel table/autofilter metadata:
    # Google Sheets supplies its own filtering UI after import. Validate the
    # fixed input/calculation extent directly instead.
    assert expenses.max_row == 204 and expenses.max_column == 26
    assert all(formula(expenses.cell(row, 26)) for row in range(5, 205))
    assert all("$Z" in formula(expenses.cell(row, 20)) for row in range(5, 205))
    assert all("$T" not in formula(expenses.cell(row, 26)) for row in range(5, 205))
    assert "SUBSTITUTE" in formula(expenses["Z5"])
    assert "ISNUMBER" in formula(expenses["Z5"])
    assert 'IF(OR($D5="",NOT(ISNUMBER($D5)))' in formula(expenses["Z5"])
    assert "IF($D5<=0" in formula(expenses["Z5"])
    assert "IF(ROUND($D5,2)<>$D5" in formula(expenses["Z5"])
    assert "IF(NOT(ISNUMBER(M5)),TRUE" in formula(expenses["Z5"])
    assert "Custom shares must add up to the expense amount." in formula(expenses["Z5"])
    assert "QUOTIENT(ROUND($D5*100,0),COUNTIF($G5:$L5,\"Yes\"))" in formula(expenses["T5"])
    assert "MOD(ROUND($D5*100,0),COUNTIF($G5:$L5,\"Yes\"))" in formula(expenses["T5"])
    assert "COUNTIF($G5:G5,\"Yes\")" in formula(expenses["T5"])

    repayments = formulas["Repayments"]
    assert not repayments.tables
    assert all(formula(repayments.cell(row, 6)) for row in range(5, 205))
    assert "SUBSTITUTE" in formula(repayments["F5"])
    assert "two decimal places" in formula(repayments["F5"])
    assert 'IF(OR($D5="",NOT(ISNUMBER($D5)))' in formula(repayments["F5"])
    assert "IF($D5<=0" in formula(repayments["F5"])

    summary = formulas["Summary"]
    assert "Settings'!$G$5" in formula(summary["B17"])
    assert "Settings'!$G$5" in formula(summary["G5"])
    assert "Fix the highlighted inputs before using the summary" in formula(summary["B17"])
    assert "<=0.01" not in formula(summary["H5"])
    assert "ROUND(G5,2)=0" in formula(summary["H5"])
    assert "<0" in formula(summary["K22"])
    assert all(formula(summary.cell(row, column)) for row in range(23, 28) for column in range(1, 4))
    assert values["Summary"]["A21"].value == "Suggested transfers that settle the current balances. Check the entries before sending money."

    settings = formulas["Settings"]
    assert "Input readiness" == values["Settings"]["F4"].value
    assert "COUNTIFS('Expenses'" in formula(settings["G5"])

    validate_equal_share_invariants()
    print("validated Google Sheets import workbook structure, direct-range formulas, readiness gate, and cent settlement contract")


if __name__ == "__main__":
    main()
