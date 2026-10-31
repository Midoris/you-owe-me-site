#!/usr/bin/env python3
"""Read-only preservation checks for the reviewed localized roommate files."""
import csv
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
COPY = json.loads((ROOT / "content/roommate-workbooks/fr.json").read_text())
NAMES = COPY["sheet_names"]
STRINGS = COPY["strings"]


def translated(value):
    if not isinstance(value, str):
        return value
    if value.startswith("="):
        for original, localized in sorted(STRINGS.items(), key=lambda item: -len(item[0])):
            value = value.replace('"' + original + '"', '"' + localized + '"')
        for original, localized in NAMES.items():
            value = value.replace("'" + original + "'!", "'" + localized + "'!")
        return value
    return STRINGS.get(value, value)


def validate_common(path):
    book = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    assert book.sheetnames == list(NAMES.values())
    assert all(not sheet.sheet_view.rightToLeft for sheet in book)
    assert [values["Listes"].cell(row, 4).value for row in range(5, 16)] == ["USD", "EUR", "GBP", "CAD", "AUD", "NZD", "JPY", "CHF", "SGD", "THB", "Autre"]
    assert all(values["Configuration"].cell(row, 2).value in (None, "") for row in range(9, 15))
    assert [values["Configuration"].cell(row, 3).value for row in range(9, 15)] == [0] * 6
    assert values["Listes"]["D15"].value == "Autre"
    assert [values["Exemple"].cell(row, 1).value for row in range(16, 19)] == ["Camille", "Hugo", "Léa"]
    assert [values["Exemple"].cell(row, 6).value for row in range(16, 19)] == [940, -400, -540]
    assert values["Exemple"]["B20"].value == 1980
    assert values["Exemple"]["B21"].value == 660
    assert [values["Exemple"].cell(row, 3).value for row in (25, 26)] == [400, 540]
    assert values["Exemple"]["F24"].value == 0
    for sheet_name in ("Dépenses", "Remboursements"):
        assert book[sheet_name].max_row == 204
        assert book[sheet_name]["A5"].number_format == "dd/mm/yyyy"
    for row in (6, 7, 8, 12):
        assert book["Exemple"].cell(row, 1).number_format == "dd/mm/yyyy"
    for row, date in [(6, "2026-05-01"), (7, "2026-05-05"), (8, "2026-05-12"), (12, "2026-05-20")]:
        assert values["Exemple"].cell(row, 1).value.date().isoformat() == date
    assert len(book.defined_names) == 5
    assert book["Dépenses"].auto_filter.ref == "A4:Z204"
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "e", (sheet.title, cell.coordinate, cell.value)
    return book, values


def validate_excel():
    path = ROOT / "downloads/fr/tableau-depenses-colocation.xlsx"
    book, values = validate_common(path)
    original = load_workbook(ROOT / "downloads/roommate-expense-tracker-template.xlsx", data_only=False)
    compared_cells = compared_formulas = 0
    for old in original:
        current = book[NAMES[old.title]]
        assert current.max_row == old.max_row and current.max_column == old.max_column
        assert str(current.merged_cells) == str(old.merged_cells)
        assert current.freeze_panes == old.freeze_panes
        assert set(current.tables) == set(old.tables)
        assert len(current.data_validations.dataValidation) == len(old.data_validations.dataValidation)
        for before, after in zip(old.data_validations.dataValidation, current.data_validations.dataValidation):
            assert before.sqref == after.sqref
            assert before.type == after.type and before.formula1 == after.formula1
        for row in old:
            for cell in row:
                actual = current[cell.coordinate]
                expected = COPY["cell_overrides"].get(f"{old.title}!{cell.coordinate}", translated(cell.value))
                assert actual.value == expected, (old.title, cell.coordinate, cell.value, expected, actual.value)
                compared_cells += 1
                compared_formulas += cell.data_type == "f"
                expected_format = "dd/mm/yyyy" if cell.number_format == "yyyy-mm-dd" else cell.number_format
                assert actual.number_format == expected_format
        for key, dimensions in old.column_dimensions.items():
            assert current.column_dimensions[key].hidden == dimensions.hidden
            assert current.column_dimensions[key].width == dimensions.width
        assert len(old.conditional_formatting) == len(current.conditional_formatting)
        for old_range, new_range in zip(old.conditional_formatting, current.conditional_formatting):
            assert old_range.sqref == new_range.sqref
            old_rules = old.conditional_formatting[old_range]
            new_rules = current.conditional_formatting[new_range]
            assert len(old_rules) == len(new_rules)
            for before, after in zip(old_rules, new_rules):
                assert after.type == before.type
                assert after.text == STRINGS.get(before.text, before.text)
                assert after.formula == [translated("=" + formula)[1:] for formula in before.formula]
    for name, definition in original.defined_names.items():
        expected = definition.attr_text
        for original_name, localized in NAMES.items():
            expected = expected.replace("'" + original_name + "'!", "'" + localized + "'!")
        assert book.defined_names[name].attr_text == expected
    with (ROOT / "downloads/fr/tableau-depenses-colocation.csv").open(newline="", encoding="utf-8") as file:
        rows = list(csv.reader(file))
    assert rows == COPY["csv"] and all(len(row) == 11 for row in rows)
    return {"compared_cells": compared_cells, "compared_formulas": compared_formulas, "native_features_preserved": True, "csv_rows": len(rows)}


def validate_sheets(path):
    book, values = validate_common(path)
    assert not book["Remboursements"].tables
    assert STRINGS["Roommate Expense Tracker — Google Sheets"] == values["Démarrage"]["A1"].value
    assert 'HYPERLINK("https://you-owe-me.com/fr/solutions/application-suivi-argent-du/","Suivez les montants restants sur iPhone")' in book["Démarrage"]["A23"].value
    assert values["Démarrage"]["A24"].value == STRINGS["Start a separate record in You Owe Me; this spreadsheet is not imported."]
    setup = book["Configuration"]["A17"].value
    for fragment in ('TRIM', 'SUBSTITUTE(B9,"  "," ")', 'NOT(ISNUMBER(C9))', 'ROUND(C16,2)<>0', STRINGS["Ready: setup inputs are valid."]):
        assert fragment in setup
    for row in range(5, 205):
        share = book["Dépenses"].cell(row, 20).value
        check = book["Dépenses"].cell(row, 26).value
        assert f'QUOTIENT(ROUND($D{row}*100,0),COUNTIF($G{row}:$L{row},"Oui"))' in share
        assert f'MOD(ROUND($D{row}*100,0),COUNTIF($G{row}:$L{row},"Oui"))' in share
        assert "SUBSTITUTE" in check and "ISNUMBER" in check
        assert STRINGS["Custom shares must add up to the expense amount."] in check
        assert "SUBSTITUTE" in book["Remboursements"].cell(row, 6).value
    assert "Listes'!$G$5" in book["Bilan"]["B17"].value
    assert "Listes'!$G$5" in book["Bilan"]["G5"].value
    assert "ROUND(G5,2)=0" in book["Bilan"]["H5"].value
    assert "<0" in book["Bilan"]["K22"].value
    assert "COUNTIFS('Dépenses'" in book["Listes"]["G5"].value
    return {"direct_range_edition": True, "readiness_gate": True, "integer_cent_allocation": True, "native_engine_pending": True}


if __name__ == "__main__":
    result = {"excel": validate_excel()}
    if len(sys.argv) == 2:
        result["sheets_intermediary"] = validate_sheets(Path(sys.argv[1]))
    print(json.dumps(result, ensure_ascii=False))
